import streamlit as st
import pandas as pd
import xml.etree.ElementTree as ET
from collections import defaultdict
from openpyxl.styles import PatternFill
from openpyxl.styles import Font
import re
from openpyxl import load_workbook
import io
import zipfile
import requests
import toml
import time
import dropbox
from dropbox.files import WriteMode

st.set_page_config(layout="wide")
st.title("🏈 Football XML Stats Parser")
st.markdown("Enter your Dropbox ZIP URL, upload the Excel template and master schedule, then generate reports for all teams.")

# Load the Dropbox token from secrets_folder/secrets.toml
secrets = toml.load("/Users/jasonwang/Desktop/CollegePressbox/secrets_folder/secrets.toml")
#DROPBOX_ACCESS_TOKEN = secrets["DROPBOX_ACCESS_TOKEN"]

# --- TEAM MAPPING (Canonical Names for Schedule/Selection) ---
TEAM_NAME_MAP = {
    "airforce": "Air Force", "akron": "Akron", "alabama": "Alabama", "appstate": "App State",
    "arizona": "Arizona", "asu": "Arizona State", "arkansas": "Arkansas", "arkstate": "Arkansas State", "astate": "Arkansas State",
    "army": "Army West Point", "auburn": "Auburn", "ballstate": "Ball State", "baylor": "Baylor",
    "boisestate": "Boise State", "bc": "Boston College", "bgsu": "Bowling Green", "buffalo": "Buffalo",
    "byu": "BYU", "cal": "California", "cmu": "Central Michigan", "charlotte": "Charlotte",
    "cincinnati": "Cincinnati", "clemson": "Clemson", "coastal": "Coastal Carolina", "colorado": "Colorado",
    "csu": "Colorado State", "delaware": "Delaware", "duke": "Duke", "ecu": "East Carolina",
    "emu": "Eastern Michigan", "fiu": "FIU", "florida": "Florida", "fau": "Florida Atlantic",
    "fsu": "Florida State", "fresnostate": "Fresno State", "georgia": "Georgia", "gasouthern": "Georgia Southern",
    "gastate": "Georgia State", "gatech": "Georgia Tech", "hawaii": "Hawaii", "houston": "Houston",
    "illinois": "Illinois", "indiana": "Indiana", "iowa": "Iowa", "iowastate": "Iowa State",
    "jmu": "James Madison", "jaxstate": "Jacksonville State", "kansas": "Kansas", "kstate": "Kansas State",
    "kennesaw": "Kennesaw State", "kentstate": "Kent State", "kentucky": "Kentucky", "liberty": "Liberty",
    "louisiana": "Louisiana", "latech": "Louisiana Tech", "louisville": "Louisville", "lsu": "LSU",
    "marshall": "Marshall", "maryland": "Maryland", "umass": "Massachusetts", "memphis": "Memphis",
    "miami": "Miami (FL)", "muohio": "Miami (OH)", "michigan": "Michigan", "michstate": "Michigan State",
    "mtsu": "Middle Tenn", "minnesota": "Minnesota", "mstate": "Mississippi State", "missouri": "Missouri",
    "mostate": "Missouri State", "navy": "Navy", "ncsu": "NC State", "nebraska": "Nebraska",
    "nevada": "Nevada", "unm": "New Mexico", "nmsu": "New Mexico State", "niu": "NIU",
    "unc": "North Carolina", "unt": "North Texas", "northwestern": "Northwestern", "notredame": "Notre Dame",
    "ohio": "Ohio", "ohiostate": "Ohio State", "oklahoma": "Oklahoma", "ostate": "Oklahoma State",
    "odu": "Old Dominion", "olemiss": "Ole Miss", "oregon": "Oregon", "oregonstate": "Oregon State",
    "pennstate": "Penn State", "pitt": "Pitt", "purdue": "Purdue", "rice": "Rice",
    "rutgers": "Rutgers", "shsu": "Sam Houston", "sdsu": "San Diego State", "sjsu": "San Jose State",
    "smu": "SMU", "usa": "South Alabama", "sc": "South Carolina", "usf": "South Florida",
    "usm": "Southern Miss", "stanford": "Stanford", "syracuse": "Syracuse", "tcu": "TCU",
    "temple": "Temple", "tennessee": "Tennessee", "texas": "Texas", "tamu": "Texas A&M",
    "texasstate": "Texas State", "texastech": "Texas Tech", "toledo": "Toledo", "troy": "Troy",
    "tulane": "Tulane", "tulsa": "Tulsa", "uab": "UAB", "ucf": "UCF",
    "ucla": "UCLA", "uconn": "UConn", "ulm": "ULM", "unlv": "UNLV",
    "usc": "USC", "utah": "Utah", "utahstate": "Utah State", "utep": "UTEP",
    "utsa": "UTSA", "vanderbilt": "Vanderbilt", "virginia": "Virginia", "vt": "Virginia Tech",
    "wfu": "Wake Forest", "washington": "Washington", "washstate": "Washington State", "wvu": "West Virginia",
    "wmu": "Western Michigan", "wisconsin": "Wisconsin", "wku": "WKU", "wyoming": "Wyoming"
}

# --- XML NAME ALIASES (Fixes for Inconsistent XML Names) ---
XML_NAME_ALIASES = {
    "appalachian state": "App State", 
    "central mich": "Central Michigan",
    "eastern mich": "Eastern Michigan",
    "western mich": "Western Michigan",
    "fla atlantic": "Florida Atlantic",
    "florida atlantic": "Florida Atlantic",
    "ga southern": "Georgia Southern",
    "Georgia  Southern": "Georgia Southern",
    "ga. southern": "Georgia Southern",
    "ull": "Louisiana",
    "northern illinois": "NIU",
    "#6/#5 notre dame": "Notre Dame",
    "nm state": "New Mexico State",
    "new mexico st": "New Mexico State",
    "pittsburgh": "Pitt",
    "south fla": "South Florida",
    "#18/#23 south florida": "South Florida",
    "southern california": "USC",
    "western ky": "WKU",
    "western kentucky": "WKU",
    "miami": "Miami (FL)",
    "#5 miami (fl)": "Miami (FL)",
    "#12 miami (fl)": "Miami (FL)",
    "#2/#2 miami (fl)": "Miami (FL)",
    "#4/#6 miami (fl)": "Miami (FL)",
    "#5/#6 miami (fl)": "Miami (FL)",
    "#5/#7 miami (fl)": "Miami (FL)",
    "#9/#9 miami (fl)": "Miami (FL)",
    "#10/#10 miami (fl)": "Miami (FL)",
    "hawai'i": "Hawaii",
    "jax state": "Jacksonville State"
}

TEAM_COLORS = {
    #AAC
    "Army West Point": ("000000", "FFE69D"),
    "Charlotte": ("203214", "FFBF00"),
    "East Carolina": ("7030A0", "FFFFFF"),
    "Florida Atlantic": ("305496", "FF0000"),
    "Memphis": ("0070C0", "A6A6A6"),
    "Navy": ("002060", "FFF2CC"),
    "North Texas": ("00B050", "FFFFFF"),
    "Rice": ("2C3C9B", "FFFFFF"),
    "South Florida": ("067C00", "FFF2CC"),
    "Temple": ("FF0000", "FFFFFF"),
    "Tulane": ("006D45", "000000"),
    "Tulsa": ("305496", "FFF2CC"),
    "UAB": ("375623", "FFBF00"),
    "UTSA": ("002060", "C65911"),
    #ACC
    "Boston College": ("C00000", "FFF2CC"),
    "California": ("203764", "FFC003"),
    "Clemson": ("ED7D31", "7030A0"),
    "Duke": ("0020FE", "FFFFFF"),
    "Florida State": ("C00000", "FFF2CC"),
    "Georgia Tech": ("FFC000", "FFFFFF"),
    "Louisville": ("FF0000", "FFFFFF"),
    "Miami (FL)": ("FF8231", "00B050"),
    "NC State": ("FF0000", "FFFFFF"),
    "North Carolina": ("00B0F0", "FFFFFF"),
    "Pitt": ("0015F4", "FFD966"),
    "SMU": ("FF0000", "0070C0"),
    "Stanford": ("BC0001", "FFFFFF"),
    "Syracuse": ("C65911", "1829BA"),
    "Virginia": ("C65911", "305496"),
    "Virginia Tech": ("A4001B", "FFFFFF"),
    "Wake Forest": ("806000", "000000"),
    #big10
    "Illinois": ("ED7D31", "305496"),
    "Indiana": ("FF0000", "FFFFFF"),
    "Iowa": ("FFF501", "000000"),
    "Maryland": ("FF0000", "FFFFFF"),
    "Michigan": ("FFFF00", "305496"),
    "Michigan State": ("548235", "FFFFFF"),
    "Minnesota": ("BD151C", "FFDF01"),
    "Nebraska": ("FF0000", "FFFFFF"),
    "Northwestern": ("7030A0", "FFFFFF"),
    "Ohio State": ("E10002", "C3C3C3"),
    "Oregon": ("FFFF00", "375623"),
    "Penn State": ("02009E", "FFFFFF"),
    "Purdue": ("000000", "FFF2CC"),
    "Rutgers": ("FF0000", "FFFFFF"),
    "UCLA": ("8EA9DB", "FFD966"),
    "USC": ("C00000", "F9F102"),
    "Washington": ("5C2785", "BF8F00"),
    "Wisconsin": ("FF0000", "FFFFFF"),
    #big12
    "Arizona": ("FF0000", "044583"),
    "Arizona State": ("B50002", "DBA900"),
    "Baylor": ("375641", "C8C900"),
    "BYU": ("040BC0", "FFFFFF"),
    "UCF": ("000000", "FFF2CC"),
    "Cincinnati": ("FF0000", "FFFFFF"),
    "Colorado": ("FFFAC4", "FFFFFF"),
    "Houston": ("FF0000", "000000"),
    "Iowa State": ("C00000", "FFC000"),
    "Kansas": ("FF0000", "305496"),
    "Kansas State": ("7030A0", "808080"),
    "Oklahoma State": ("C65911", "808080"),
    "TCU": ("5D2784", "FFFFFF"),
    "Texas Tech": ("FF0000", "000000"),
    "Utah": ("FF0000", "FFFFFF"),
    "West Virginia": ("203764", "FFD802"),
    #CUSA
    "Delaware": ("00B0F0", "FFFFFF"),
    "FIU": ("002993", "FFC000"),
    "Jacksonville State": ("FF0000", "FFFFFF"),
    "Kennesaw State": ("FFFF00", "000000"),
    "Louisiana Tech": ("305496", "FF0000"),
    "Liberty": ("FF0000", "305496"),
    "Middle Tenn": ("0070C0", "A6A6A6"),
    "Missouri State": ("C65911", "FFFFFF"),
    "New Mexico State": ("C93530", "FFFFFF"),
    "Sam Houston": ("C65911", "FFFFFF"),
    "UTEP": ("ED7D31", "FFFFFF"),
    "WKU": ("FF0000", "FFFFFF"),
    #IND
    "UConn": ("044583", "FFFFFF"),
    "Notre Dame": ("002060", "EBB000"),
    #MAC
    "Akron": ("012548", "806000"),
    "Ball State": ("FF0000", "FFC000"),
    "Bowling Green": ("C65911", "5D2B09"),
    "Buffalo": ("305496", "FFFFFF"),
    "Central Michigan": ("CF0002", "FFC000"),
    "Eastern Michigan": ("306805", "FFFFFF"),
    "Kent State": ("203764", "FFE509"),
    "Miami (OH)": ("FF0000", "FFFFFF"),
    "NIU": ("DB0003", "838383"),
    "Ohio": ("0A711B", "FFFFFF"),
    "Toledo": ("002060", "EFEE00"),
    "Massachusetts": ("C60E11", "FFFFFF"),
    "Western Michigan": ("806000", "000000"),
    #SunBelt
    "App State": ("FFC000", "000000"),
    "Arkansas State": ("FF0000", "000000"),
    "Costal Carolina": ("00E4D5", "000000"),
    "Georgia Southern": ("044583", "FFFFFF"),
    "Georgia State": ("2F75B5", "FFFFFF"),
    "James Madison": ("7030A0", "BF8F00"),
    "Louisiana": ("FF0000", "FFFFFF"),
    "Marshall": ("00B050", "FFFFFF"),
    "Old Dominion": ("044583", "808080"),
    "South Alabama": ("FF0000", "FFFFFF"),
    "Southern Miss": ("806000", "000000"),
    "Texas State": ("A00001", "806000"),
    "Troy": ("C00000", "A6A6A6"),
    "ULM": ("950001", "FFC000"),
    #SEC
    "Alabama": ("C80001", "FFFFFF"),
    "Arkansas": ("C80001", "FFFFFF"),
    "Auburn": ("FF8700", "002060"),
    "Florida": ("3505A0", "FF9219"),
    "Georgia": ("000000", "FF0000"),
    "Kentucky": ("305496", "FFFFFF"),
    "LSU": ("7030A0", "FFFF00"),
    "Mississippi": ("8EA9DB", "FF0000"),
    "Mississippi State": ("8E0001", "5E5E5E"),
    "Missouri": ("BF8F00", "000000"),
    "Oklahoma": ("C00000", "FFF2CC"),
    "South Carolina": ("950001", "000000"),
    "Tennessee": ("FF7217", "000000"),
    "Texas A&M": ("8A0001", "FFFFFF"),
    "Texas": ("C65911", "FFFFFF"),
    "Vanderbilt": ("BF8F00", "000000"),
    #PAC12
    "Oregon State": ("ED7D31", "000000"),
    "Washington State": ("FF0000", "BFBFBF"),
    #MTWest
    "Air Force": ("0070C0", "A6A6A6"),
    "Boise State": ("002B91", "ED7D31"),
    "Colorado State": ("365523", "FFFFFF"),
    "Fresno State": ("FF0000", "FFFFFF"),
    "Hawaii": ("048A00", "FFFFFF"),
    "Nevada": ("002060", "A6A6A6"),
    "New Mexico": ("C40002", "A6A6A6"),
    "San Diego State": ("FF0000", "262626"),
    "San Jose State": ("305496", "FFFF00"),
    "UNLV": ("C00002", "A6A6A6"),
    "Utah State": ("023668", "FFFFFF"),
    "Wyoming": ("BF8F00", "806000")
}

TEAM_HEADER_COLORS = {
    #AAC Special Case
    "East Carolina": {"A4": "FFFF00","A14": "000000"},
    "Florida Atlantic": {"A4": "FFFFFF","A14": "FFFFFF"},
    "Navy": {"A4": "FFF2CC","A14": "000000"},
    "Tulane": {"A4": "FFFFFF","A14": "FFFFFF"},
    "Tulsa": {"A4": "FFF2CC","A14": "000000"},
    #ACC Special Case
    "Boston College": {"A4": "FFF2CC","A14": "000000"},
    "Clemson": {"A4": "000000","A14": "FFFFFF"},
    "Florida State": {"A4": "FFF2CC","A14": "000000"},
    "Georgia Tech": {"A4": "000000","A14": "000000"},
    "Miami (FL)": {"A4": "FFFFFF","A14": "FFFFFF"},
    "North Carolina": {"A4": "FFFFFF","A14": "FFFFFF"},
    "SMU": {"A4": "FFFFFF","A14": "FFFFFF"},
    "Syracuse": {"A4": "FFFFFF","A14": "FFFFFF"},
    "Virginia": {"A4": "FFFFFF","A14": "FFFFFF"},
    "Wake Forest": {"A4": "FFFFFF","A14": "FFFFFF"},
    #BIG10 Special Case
    "Illinois": {"A4": "FFFFFF","A14": "FFFFFF"},
    "Iowa": {"A4": "000000","A14": "FFFFFF"},
    "Michigan": {"A4": "000000","A14": "FFFFFF"},
    "Oregon": {"A4": "000000","A14": "FFFFFF"},
    "Purdue": {"A4": "FFF2CC","A14": "000000"},
    "UCLA": {"A4": "000000","A14": "000000"},
    "Washington": {"A4": "FFFFFF","A14": "FFFFFF"},
    #BIG12 Special Case
    "Arizona": {"A4": "FFFFFF","A14": "FFFFFF"},
    "UCF": {"A4": "FFF2CC","A14": "000000"},
    "Colorado": {"A4": "000000","A14": "000000"},
    "Houston": {"A4": "FFFFFF","A14": "FFFFFF"},
    "Kansas": {"A4": "FFFFFF","A14": "FFFFFF"},
    "Kansas State": {"A4": "FFFFFF","A14": "FFFFFF"},
    "Oklahoma State": {"A4": "FFFFFF","A14": "FFFFFF"},
    "Texas Tech": {"A4": "FFFFFF","A14": "FFFFFF"},
    #CUSA Special Case
    "Kennesaw State": {"A4": "000000","A14": "000000"},
    "Louisiana Tech": {"A4": "FFFFFF","A14": "FFFFFF"},
    "Liberty": {"A4": "FFFFFF","A14": "FFFFFF"},
    #MAC Special Case
    "Akron": {"A4": "FFFFFF","A14": "FFFFFF"},
    "Bowling Green": {"A4": "FFFFFF","A14": "FFFFFF"},
    "NIU": {"A4": "FFFFFF","A14": "FFFFFF"},
    "Western Michigan": {"A4": "FFFFFF","A14": "FFFFFF"},
    #SunBelt Special Case
    "App State": {"A4": "000000","A14": "FFFFFF"},
    "Arkansas State": {"A4": "FFFFFF","A14": "FFFFFF"},
    "Coastal Carolina": {"A4": "000000","A14": "FFFFFF"},
    "James Madison": {"A4": "FFFFFF","A14": "FFFFFF"},
    "Old Dominion": {"A4": "FFFFFF","A14": "FFFFFF"},
    "Southern Miss": {"A4": "FFFFFF","A14": "FFFFFF"},
    "Texas State": {"A4": "FFFFFF","A14": "FFFFFF"},
    "Troy": {"A4": "FFFFFF","A14": "FFFFFF"},
    #SEC Special Case
    "Auburn": {"A4": "FFFFFF","A14": "FFFFFF"},
    "Georgia": {"A4": "FFFFFF","A14": "FFFFFF"},
    "Mississippi": {"A4": "000000","A14": "FFFFFF"},
    "Mississippi State": {"A4": "FFFFFF","A14": "FFFFFF"},
    "Missouri": {"A4": "000000","A14": "FFFFFF"},
    "South Carolina": {"A4": "FFFFFF","A14": "FFFFFF"},
    "Tennessee": {"A4": "FFFFFF","A14": "FFFFFF"},
    "Vanderbilt": {"A4": "FFFFFF","A14": "FFFFFF"},
    #PAC12 Special Case
    "Oregon State": {"A4": "000000","A14": "FFFFFF"},
    #MTWest Special Case
    "San Diego State": {"A4": "FFFFFF","A14": "FFFFFF"},
    "Wyoming": {"A4": "FFFFFF","A14": "FFFFFF"},
}

# NEW: Function to get access token using refresh token
def get_dropbox_client():
    """
    Creates a Dropbox client using refresh token for long-term authentication.
    This way you don't need to regenerate access tokens.
    """
    try:
        # Check if using refresh token (recommended)
        if "DROPBOX_REFRESH_TOKEN" in secrets:
            return dropbox.Dropbox(
                app_key=secrets["DROPBOX_APP_KEY"],
                app_secret=secrets["DROPBOX_APP_SECRET"],
                oauth2_refresh_token=secrets["DROPBOX_REFRESH_TOKEN"]
            )
        # Fallback to old method with access token
        elif "DROPBOX_ACCESS_TOKEN" in secrets:
            return dropbox.Dropbox(secrets["DROPBOX_ACCESS_TOKEN"])
        else:
            st.error("No Dropbox credentials found in secrets.toml")
            return None
    except Exception as e:
        st.error(f"Error creating Dropbox client: {e}")
        return None

def apply_team_colors(ws, primary_hex, secondary_hex):
    """Apply primary and secondary colors to fixed ranges in Data Output sheet."""
    primary_fill = PatternFill(start_color=primary_hex, end_color=primary_hex, fill_type="solid")
    secondary_fill = PatternFill(start_color=secondary_hex, end_color=secondary_hex, fill_type="solid")

    # Primary color ranges
    for row in range(4, 14):  # B4:B13
        ws[f"B{row}"].fill = primary_fill
        ws[f"A{row}"].fill = primary_fill
    for row in range(29, 49):  # B29:B48
        ws[f"B{row}"].fill = primary_fill
        ws[f"A{row}"].fill = primary_fill

    # Secondary color ranges
    for row in range(14, 29):  # B14:B28
        ws[f"B{row}"].fill = secondary_fill
        ws[f"A{row}"].fill = secondary_fill

def apply_team_header_colors(ws, core_team_name):
    """Apply team-specific *font* colors to merged header cells (A4=A29)."""
    if core_team_name not in TEAM_HEADER_COLORS:
        return  # skip if not found

    colors = TEAM_HEADER_COLORS[core_team_name]

    def set_font_color(cell_addr, hex_color):
        ws[cell_addr].font = Font(name = "Arial", size = 14, bold = True, color=f"FF{hex_color}")

    # Apply font colors (A4 = A29)
    set_font_color("A4", colors["A4"])
    set_font_color("A29", colors["A4"])
    set_font_color("A14", colors["A14"])


# --- HELPER FUNCTIONS ---

def normalize_team_name(name):
    """
    Normalizes team names from XML, aggressively handling inconsistent spacing and aliases.
    """
    if not name:
        return ""
        
    # 1. AGGRESSIVE WHITESPACE CLEANUP
    name_cleaned = re.sub(r'\s+', ' ', name).strip()
    
    # 2. Apply existing cleanup rules
    name_cleaned = name_cleaned.replace('.', '')
    name_cleaned = re.sub(r'^#\d+\s+', '', name_cleaned)
    name_cleaned = re.sub(r'\bSt\b', 'State', name_cleaned)
    name_cleaned = name_cleaned.strip()
    
    # 3. Alias Check
    name_lower = name_cleaned.lower()
    
    if name_lower in XML_NAME_ALIASES:
        return XML_NAME_ALIASES[name_lower]
    
    # 4. Default Return
    return name_cleaned

def normalize_schedule_team_name(name):
    """
    Extra normalization for team names in the schedule file to match our canonical names.
    """
    if not name or pd.isna(name):
        return ""
    
    name = str(name).strip()
    
    # Common schedule variations to canonical names
    schedule_aliases = {
        "army": "Army West Point",
        "middle tennessee": "Middle Tenn",
        "mt": "Middle Tenn",
        "mtsu": "Middle Tenn",
        "central florida": "UCF",
        "connecticut": "UConn",
        "louisiana monroe": "ULM",
        "ul monroe": "ULM",
        "miami (florida)": "Miami (FL)",
        "miami fl": "Miami (FL)",
        "miami-fl": "Miami (FL)",
        "miami (ohio)": "Miami (OH)",
        "miami oh": "Miami (OH)",
        "miami-oh": "Miami (OH)",
    }
    
    name_lower = name.lower()
    
    # Check schedule-specific aliases first
    if name_lower in schedule_aliases:
        return schedule_aliases[name_lower]
    
    # Then use the regular normalize function
    return normalize_team_name(name)

def load_master_schedule(schedule_file):
    """
    Load the master schedule Excel file and return a dictionary mapping team names to their games.
    Returns: dict where key = team_name, value = list of (date, opponent, is_home) tuples
    """
    try:
        df = pd.read_excel(schedule_file, engine='openpyxl')
        
        # Verify required columns exist
        if 'Date' not in df.columns or 'Visitor' not in df.columns or 'Home' not in df.columns:
            st.error("Master schedule must have columns: 'Date' (column B), 'Visitor' (column D), and 'Home' (column F)")
            return None
        
        # Build schedule dictionary
        team_schedules = defaultdict(list)
        
        for _, row in df.iterrows():
            date_val = row['Date']
            visitor = normalize_schedule_team_name(row['Visitor'])
            home = normalize_schedule_team_name(row['Home'])
            
            # Skip if either team name is empty
            if not visitor or not home:
                continue
            
            # Parse date
            if pd.isna(date_val):
                continue
            
            if isinstance(date_val, str):
                game_date = pd.to_datetime(date_val, errors='coerce')
            else:
                game_date = pd.to_datetime(date_val)
            
            if pd.isna(game_date):
                continue
            
            # Add game to visitor's schedule (away game)
            team_schedules[visitor].append((game_date, home, False))
            
            # Add game to home team's schedule (home game)
            team_schedules[home].append((game_date, visitor, True))
        
        # Sort each team's schedule by date
        for team in team_schedules:
            team_schedules[team].sort(key=lambda x: x[0])
        
        return dict(team_schedules)
        
    except Exception as e:
        st.error(f"Error loading master schedule: {e}")
        return None

def build_cumulative_df(core_stats, master_order, category):
    data = []
    stat_keys = {
        'rush': ['Rush Att','Rush Yds','Rush TD'],
        'pass': ['Pass Comp','Pass Att','Pass Yds','Pass TD','Pass INT'],
        'recv': ['Rec','Rec Yds','Rec TD']
    }
    
    for p in master_order[category]:
        s = core_stats.get(p, {k: 0 for k in stat_keys['rush'] + stat_keys['pass'] + stat_keys['recv']})

        if category == 'rush':
            row = {'Player': p, 'Att': s['Rush Att'], 'Yds': s['Rush Yds'], 'TD': s['Rush TD']}
        elif category == 'pass':
            row = {'Player': p, 'Comp': s['Pass Comp'], 'Att': s['Pass Att'], 'Yds': s['Pass Yds'], 'TD': s['Pass TD'], 'INT': s['Pass INT']}
        elif category == 'recv':
            row = {'Player': p, 'Rec': s['Rec'], 'Yds': s['Rec Yds'], 'TD': s['Rec TD']}
        data.append(row)
    return pd.DataFrame(data)

def fill_excel_df(ws, df, start_row, columns):
    for i, row in df.iterrows():
        excel_row = start_row + i
        for col_letter, col_name in columns.items():
            ws[f"{col_letter}{excel_row}"] = row[col_name]

def set_cell_value_safe(ws, row, col_letter, value):
    cell_coord = f"{col_letter}{row}"
    for merged in ws.merged_cells.ranges:
        if cell_coord in merged:
            if cell_coord != merged.start_cell.coordinate:
                return
    ws[cell_coord].value = value

def build_game_df(core_game_stats, master_list, category):
    data = []
    stat_keys = {
        'rush': ['Rush Att', 'Rush Yds', 'Rush TD'],
        'pass': ['Pass Comp', 'Pass Att', 'Pass Yds', 'Pass TD', 'Pass INT', 'Sacks'],
        'recv': ['Rec', 'Rec Yds', 'Rec TD']
    }

    for p in master_list:
        stats = core_game_stats.get(p, {k: 0 for k in stat_keys['rush'] + stat_keys['pass'] + stat_keys['recv']})
        
        if category == 'rush':
            row = {'Player': p, 'Att': stats['Rush Att'], 'Yds': stats['Rush Yds'], 'TD': stats['Rush TD']}
        elif category == 'pass':
            row = {'Player': p, 'Comp': stats['Pass Comp'], 'Att': stats['Pass Att'],
                   'Yds': stats['Pass Yds'], 'TD': stats['Pass TD'], 'INT': stats['Pass INT']}
        elif category == 'recv':
            row = {'Player': p, 'Rec': stats['Rec'], 'Yds': stats['Rec Yds'], 'TD': stats['Rec TD']}
        data.append(row)
    return pd.DataFrame(data)


# --- CORE PROCESSING FUNCTION ---
def process_team(core_team_name, team_stats, game_stats, parsed_games, template_data, keep_vba, master_schedule, progress_bar, progress_index, total_teams):
    
    # Update progress bar
    progress_percent = (progress_index / total_teams)
    progress_bar.progress(progress_percent, text=f"Processing {core_team_name}... ({progress_index}/{total_teams})")
    
    core_stats = team_stats[core_team_name]

    # --- Get schedule from master schedule file ---
    if core_team_name not in master_schedule:
        st.warning(f"Team {core_team_name} not found in master schedule. Skipping.")
        return None
    
    schedule_games = master_schedule[core_team_name]
    
    if not schedule_games:
        return None

    # --- Build master order ---
    master_order = {'rush': [], 'pass': [], 'recv': []}

    for game_idx, (game_date, opponent, is_home) in enumerate(schedule_games, start=1):
        # Find matching game in game_stats
        game_found = False
        for game_name in game_stats:
            if core_team_name in game_stats[game_name]:
                # Check if opponent is in this game
                teams_in_game = list(game_stats[game_name].keys())
                if opponent in teams_in_game or any(normalize_team_name(t) == opponent for t in teams_in_game):
                    core_game_stats = game_stats[game_name][core_team_name]
                    game_found = True
                    break
        
        if not game_found:
            continue

        rush_players = [p for p, s in core_game_stats.items() if s['Rush Att'] > 0 or s['Rush Yds'] != 0 or s['Rush TD'] != 0]
        pass_players = [p for p, s in core_game_stats.items() if (s['Pass Att'] > 0 or s['Pass Yds'] != 0 or s['Pass TD'] != 0 or s['Pass Comp'] > 0 or s['Pass INT'] > 0 or s['Sacks'] > 0)]
        recv_players = [p for p, s in core_game_stats.items() if s['Rec'] > 0 or s['Rec Yds'] != 0 or s['Rec TD'] != 0]

        if game_idx == 1:
            rush_players.sort(key=lambda p: core_game_stats[p]['Rush Yds'], reverse=True)
            pass_players.sort(key=lambda p: core_game_stats[p]['Pass Yds'], reverse=True)
            recv_players.sort(key=lambda p: core_game_stats[p]['Rec Yds'], reverse=True)

        for category, players in [('rush', rush_players), ('pass', pass_players), ('recv', recv_players)]:
            for p in players:
                if p not in master_order[category]:
                    master_order[category].append(p)

    # --- Build cumulative DataFrames ---
    rush_df = build_cumulative_df(core_stats, master_order, 'rush')
    pass_df = build_cumulative_df(core_stats, master_order, 'pass')
    recv_df = build_cumulative_df(core_stats, master_order, 'recv')
    
    # --- Load Excel template ---
    template_buffer = io.BytesIO(template_data) 
    
    try:
        wb = load_workbook(template_buffer, keep_vba=keep_vba)
        ws = wb["BOX SCORES"]
        ws_sorting = wb["Sorting Sheet"] if "Sorting Sheet" in wb.sheetnames else None
    except KeyError:
        st.error(f"Failed to find the worksheet named 'BOX SCORES' in the template for {core_team_name}. Skipping.")
        return None
    except Exception as e:
        st.error(f"Error loading Excel template for {core_team_name}: {e}. Skipping.")
        return None

    # Fill cumulative stats
    fill_excel_df(ws, rush_df, 3, {'B':'Player','C':'Att','D':'Yds','F':'TD'})
    fill_excel_df(ws, pass_df, 3, {'G':'Player','H':'Comp','I':'Att','K':'Yds','M':'TD','O':'INT'})
    fill_excel_df(ws, recv_df, 3, {'R':'Player','S':'Rec','T':'Yds','V':'TD'})

    # --- Fill game stats ---
    for game_idx, (game_date, opponent, is_home) in enumerate(schedule_games, start=1):
        # Find matching game in game_stats
        game_found = False
        for game_name in game_stats:
            if core_team_name in game_stats[game_name]:
                teams_in_game = list(game_stats[game_name].keys())
                if opponent in teams_in_game or any(normalize_team_name(t) == opponent for t in teams_in_game):
                    core_game_stats = game_stats[game_name][core_team_name]
                    game_found = True
                    break
        
        if not game_found:
            continue
        
        # Build game DFs using the master order
        rush_df_game = build_game_df(core_game_stats, master_order['rush'], 'rush')
        pass_df_game = build_game_df(core_game_stats, master_order['pass'], 'pass')
        recv_df_game = build_game_df(core_game_stats, master_order['recv'], 'recv')

        # Fill game blocks
        base_row = 32 + (game_idx - 1) * 29
        fill_excel_df(ws, rush_df_game, base_row, {'C': 'Att', 'D': 'Yds', 'F': 'TD'})
        fill_excel_df(ws, pass_df_game, base_row, {'H': 'Comp', 'I': 'Att', 'K': 'Yds', 'M': 'TD', 'O': 'INT'})
        fill_excel_df(ws, recv_df_game, base_row, {'S': 'Rec', 'T': 'Yds', 'V': 'TD'})

        # Fill player names
        for i, p_name in enumerate(master_order['rush']):
             set_cell_value_safe(ws, base_row + i, 'B', p_name)
        
        for i, p_name in enumerate(master_order['pass']):
            set_cell_value_safe(ws, base_row + i, 'G', p_name)

        for i, p_name in enumerate(master_order['recv']):
            set_cell_value_safe(ws, base_row + i, 'R', p_name)

    # --- Fill Schedule Columns (Z, AA, AB) in BOX SCORES ---
    for i, (game_date, opponent, is_home) in enumerate(schedule_games[:17], start=2):
        # Format date as M/D
        ws[f"Z{i}"] = game_date.strftime("%-m/%-d") if game_date.strftime("%m")[0] != "0" else game_date.strftime("%m/%d").lstrip("0").replace("/0", "/")
        
        # Determine if playoff game
        is_playoff = (i - 1) > 12
        
        # Fill AA column (at/home indicator)
        if is_playoff and not is_home:
            ws[f"AA{i}"] = "†"
        elif not is_home:
            ws[f"AA{i}"] = "at"
        else:
            ws[f"AA{i}"] = ""
        
        # Fill AB column (opponent name)
        ws[f"AB{i}"] = opponent

    # --- Fill Schedule Columns in Sorting Sheet (if exists) ---
    if ws_sorting:
        for i, (game_date, opponent, is_home) in enumerate(schedule_games[:17], start=2):
            ws_sorting[f"Z{i}"] = game_date.strftime("%-m/%-d") if game_date.strftime("%m")[0] != "0" else game_date.strftime("%m/%d").lstrip("0").replace("/0", "/")
            
            is_playoff = (i - 1) > 12
            
            if is_playoff and not is_home:
                ws_sorting[f"AA{i}"] = "†"
            elif not is_home:
                ws_sorting[f"AA{i}"] = "at"
            else:
                ws_sorting[f"AA{i}"] = ""
            
            ws_sorting[f"AB{i}"] = opponent

    # Apply team colors
    if "Data Output" in wb.sheetnames and core_team_name in TEAM_COLORS:
        ws_colors = wb["Data Output"]
        primary_hex, secondary_hex = TEAM_COLORS[core_team_name]
        apply_team_colors(ws_colors, primary_hex, secondary_hex)
        apply_team_header_colors(ws_colors, core_team_name)

    # --- Save and Return file data ---
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    special_cases = {
        "Army West Point": "Army_Ind",
        "Miami (FL)": "Miami_Fl_Ind",
        "Miami (OH)": "Miami_Ohio_Ind",
        "Georgia State": "Georgia_State_Ind",
        "Jacksonville State": "JacksonvilleSt_Ind",
        "Louisiana Tech": "La_Tech_Ind",
        "Missouri State": "MissouriSt_Ind",
        "Sam Houston": "SamHoustonSt_Ind"
    }

    if core_team_name in special_cases:
        file_name = special_cases[core_team_name]
    else:
        file_name = core_team_name.replace(" ", "_")
        if "State" in core_team_name and "Georgia State" not in core_team_name:
            file_name = file_name.replace("State", "St")
        file_name += "_Ind"

    file_name += ".xlsm" if keep_vba else ".xlsx"

    return output, file_name


# ----------------------------------------------------------------------
# --- MAIN APPLICATION LOGIC ---
# ----------------------------------------------------------------------

dropbox_zip_url = st.text_input(
    "Enter the Direct ZIP Download Link for the XML files (must end with &dl=1):",
    "" 
)

template_file = st.file_uploader("Upload your Excel template (.xlsx or .xlsm)", type=["xlsx", "xlsm"])

schedule_file = st.file_uploader("Upload the Master Schedule Excel file", type=["xlsx", "xlsm"])

# Only start the process if the button is clicked and all files are provided
if st.button("Generate Reports") and dropbox_zip_url and template_file and schedule_file:
    
    # --- Load Master Schedule First ---
    with st.spinner("Loading master schedule..."):
        master_schedule = load_master_schedule(schedule_file)
        if not master_schedule:
            st.error("Failed to load master schedule. Please check the file format.")
            st.stop()
        st.success(f"✅ Master schedule loaded for {len(master_schedule)} teams")
    
    # --- Download and Extract XML Files ---
    try:
        with st.spinner(f"Downloading files from {dropbox_zip_url}..."):
            response = requests.get(dropbox_zip_url, stream=True, timeout=60)
            response.raise_for_status() 
            zip_buffer = io.BytesIO(response.content)

        all_xml_data = {}  
        with st.spinner("Extracting and preparing XML files..."):
            with zipfile.ZipFile(zip_buffer, 'r') as z:
                for file_info in z.infolist():
                    file_name = file_info.filename
                    if file_name.endswith('.xml') and not file_info.is_dir() and not file_name.startswith('__MACOSX'):
                        content = z.read(file_name)
                        all_xml_data[file_name] = io.BytesIO(content)
    
    except requests.exceptions.RequestException as e:
        st.error(f"Error downloading ZIP file. Please ensure the link is correct and ends with `&dl=1`. Error: {e}")
        st.stop()
    except zipfile.BadZipFile:
        st.error("The downloaded file is not a valid ZIP archive. Ensure the URL ends with `&dl=1`.")
        st.stop()
    except Exception as e:
        st.error(f"An unexpected error occurred during file extraction: {e}")
        st.stop()

    # --- Parse XML Filenames ---
    parsed_games = []
    
    for file_name_with_path, file_stream in all_xml_data.items():
        filename_base_with_ext = file_name_with_path.split('/')[-1].split('\\')[-1]
        filename_base = filename_base_with_ext.replace(".xml", "")
        
        match = re.match(r"(\d{6})gamebook_(.+)_(.+)", filename_base)
        
        if match:
            date_str, team1, team2 = match.groups()
            game_date = pd.to_datetime(date_str, format="%y%m%d")
            parsed_games.append((game_date, team1.lower(), team2.lower(), file_stream, filename_base_with_ext))
        else:
            st.warning(f"Filename format not recognized: {file_name_with_path}. Skipping.")

    if not parsed_games:
        st.error("No valid XML files found in the ZIP archive matching the required naming convention.")
        st.stop()
        
    parsed_games.sort(key=lambda x: x[0]) 

    # --- Parse XML Data ---
    team_stats = defaultdict(lambda: defaultdict(lambda: {
        'Rush Att': 0, 'Rush Yds': 0, 'Rush TD': 0, 'Rec': 0, 'Rec Yds': 0, 'Rec TD': 0,
        'Pass Comp': 0, 'Pass Att': 0, 'Pass Yds': 0, 'Pass TD': 0, 'Pass INT': 0, 'Sacks': 0
    }))
    team_names_set = set()
    game_stats = {}

    with st.spinner("Parsing XML data and aggregating stats..."):
        for game_date, team1, team2, file_stream, file_name in parsed_games: 
            file_stream.seek(0)
            
            try:
                tree = ET.parse(file_stream) 
                root = tree.getroot()
            except ET.ParseError:
                st.error(f"Error parsing XML file: {file_name}. Skipping.")
                continue

            game_name = file_name.replace(".xml", "")
            game_stats[game_name] = {}

            for team in root.findall('.//team'):
                team_name = normalize_team_name(team.attrib.get('name', ''))
                team_names_set.add(team_name)
                players_dict = {}

                for player in team.findall('./player'):
                    name = player.attrib.get('name', '')
                    opos = player.attrib.get('opos', None)
                    if opos in (None, 'None') and name != "TEAM":
                        continue
                    if ',' in name:
                        last, first = name.split(',', 1)
                        name = f"{first.strip()} {last.strip()}"

                    stats = {k: 0 for k in team_stats[team_name][name].keys()}

                    if (r := player.find('rush')) is not None:
                        stats['Rush Att'] = int(r.attrib.get('att', 0))
                        stats['Rush Yds'] = int(r.attrib.get('yds', 0))
                        stats['Rush TD'] = int(r.attrib.get('td', 0))

                    if (rcv := player.find('rcv')) is not None:
                        stats['Rec'] = int(rcv.attrib.get('no', 0))
                        stats['Rec Yds'] = int(rcv.attrib.get('yds', 0))
                        stats['Rec TD'] = int(rcv.attrib.get('td', 0))

                    if (p := player.find('pass')) is not None:
                        stats['Pass Comp'] = int(p.attrib.get('comp', 0))
                        stats['Pass Att'] = int(p.attrib.get('att', 0))
                        stats['Pass Yds'] = int(p.attrib.get('yds', 0))
                        stats['Pass TD'] = int(p.attrib.get('td', 0))
                        stats['Pass INT'] = int(p.attrib.get('int', 0))
                        stats['Sacks'] = int(p.attrib.get('sacks', 0))
                    else:
                        stats['Sacks'] = 0
                    
                    for k in stats:
                        team_stats[team_name][name][k] += stats[k]
                    players_dict[name] = stats

                game_stats[game_name][team_name] = players_dict

    # --- Process Teams ---
    valid_teams = sorted([team for team in team_names_set if team in TEAM_NAME_MAP.values()])
    total_teams = len(valid_teams)

    if not valid_teams:
        st.error("No teams found in the XML files that match the accepted team list for report generation.")
        st.stop()
        
    keep_vba = template_file.name.endswith(".xlsm")
    template_data = template_file.getvalue() 
    
    st.header(f"📊 Starting Automated Report Generation for {total_teams} Teams...")
    progress_bar = st.progress(0, text="Initializing...")

    results = []
    
    for idx, team in enumerate(valid_teams):
        result = process_team(team, team_stats, game_stats, parsed_games, template_data, keep_vba, master_schedule, progress_bar, idx + 1, total_teams)
        
        if result:
            results.append(result)
        else:
            st.warning(f"Could not generate report for {team}.")

    progress_bar.progress(1.0, text="✅ All reports generated!")
    
    # --- Create ZIP ---
    st.header("📦 Packaging All Sheets into One File")

    if results:
        zip_buffer = io.BytesIO()
        
        with st.spinner("Creating ZIP file..."):
            with zipfile.ZipFile(zip_buffer, 'w', zipfile.ZIP_DEFLATED) as zip_file:
                for file_data_io, file_name in results:
                    zip_file.writestr(file_name, file_data_io.getvalue())
        
        zip_buffer.seek(0)

        st.download_button(
            label = "📥 Download All Reports as ZIP",
            data = zip_buffer,
            file_name = "All_Team_Sheets.zip",
            mime = "application/zip"
        )

        # --- Dropbox Upload ---

        dbx = get_dropbox_client()
        if not dbx:
            st.error("Failed to create Dropbox client")
        else:
            dropbox_folder = "/FootballReports"
            file_path = f"{dropbox_folder}/All_Team_Sheets.zip"

            with st.spinner("Uploading ZIP to Dropbox..."):
                try:
                    zip_buffer.seek(0)
                    dbx.files_upload(zip_buffer.read(), file_path, mode=dropbox.files.WriteMode.overwrite)
                    st.success(f"✅ All reports uploaded to Dropbox at {file_path}!")
                except Exception as e:
                    st.error(f"Failed to upload ZIP to Dropbox: {e}")

    else:
        st.error("No reports were successfully processed. Check the warnings above for potential errors.")
